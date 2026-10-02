#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
excel_a_json.py — Regenera data/alimentos.json desde COSECHA_Base_Alimentos.xlsx.

Lee el libro (hojas 'Alimentos' y 'Porciones'), VALIDA esquema y tipos y
escribe un alimentos.json con la MISMA estructura que el generado por
scripts/generar_base.py. El campo kcalDesvia no vive en el Excel: se
recalcula aquí con la misma regla (|kcal − (4P+4C+9G)| / max(kcal,1) > 0.15)
y solo MARCA, nunca corrige.

Reglas que hace cumplir (si una falla, aborta con el detalle):
  - encabezados exactos en 'Alimentos' y 'Porciones';
  - id con forma ALI-0000, único;
  - nombre, grupo, estado, fuente y licencia no vacíos;
  - kcal, proteina_g, carbohidratos_g y grasa_g numéricos y >= 0 (NUNCA vacíos);
  - fibra_g numérica >= 0 o vacía (vacía = la fuente no declara el dato);
  - clave (nombre normalizado sin acentos + estado) única: 0 duplicados;
  - cada porción apunta a un id existente, con medida y fuente no vacías
    y gramos numérico > 0.

Uso:
  python3 excel_a_json.py RUTA_AL_XLSX [--out RUTA_JSON]

  --out  default: <carpeta del xlsx>/data/alimentos.json
"""
import argparse
import json
import os
import re
import sys
import unicodedata

COLS_ALIMENTOS = ["id", "nombre", "sinonimos", "grupo", "estado", "kcal",
                  "proteina_g", "carbohidratos_g", "grasa_g", "fibra_g",
                  "fuente", "id_fuente", "licencia"]
COLS_PORCIONES = ["id", "alimento", "medida", "gramos", "fuente"]
RE_ID = re.compile(r"^ALI-\d{4}$")


def normalizar(s):
    s = unicodedata.normalize("NFD", s.lower().strip())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def fallo(msg):
    sys.exit(f"ERROR de validación: {msg}")


def texto(v, campo, fila, obligatorio=True):
    if v is None or (isinstance(v, str) and not v.strip()):
        if obligatorio:
            fallo(f"fila {fila}: '{campo}' vacío")
        return ""
    if not isinstance(v, str):
        fallo(f"fila {fila}: '{campo}' debe ser texto, es {type(v).__name__}")
    return v.strip()


def numero(v, campo, fila, obligatorio=True, minimo=0):
    if v is None or (isinstance(v, str) and not v.strip()):
        if obligatorio:
            fallo(f"fila {fila}: '{campo}' vacío (los macros nunca pueden "
                  "quedar vacíos)")
        return None
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        fallo(f"fila {fila}: '{campo}' debe ser número, es {type(v).__name__}")
    if v < minimo:
        fallo(f"fila {fila}: '{campo}'={v} fuera de rango (mínimo {minimo})")
    return v


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx", help="ruta a COSECHA_Base_Alimentos.xlsx")
    ap.add_argument("--out", default=None, help="ruta del JSON de salida")
    args = ap.parse_args()

    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("ERROR: falta openpyxl (pip install openpyxl)")

    if not os.path.isfile(args.xlsx):
        sys.exit(f"ERROR: no existe {args.xlsx}")
    wb = load_workbook(args.xlsx, data_only=True)
    for hoja in ("Alimentos", "Porciones"):
        if hoja not in wb.sheetnames:
            fallo(f"falta la hoja '{hoja}'")

    # ---------------- hoja Alimentos
    ws = wb["Alimentos"]
    filas = list(ws.iter_rows(values_only=True))
    if not filas or [c if c is not None else "" for c in filas[0]][:13] != COLS_ALIMENTOS:
        fallo(f"encabezados de 'Alimentos' distintos de {COLS_ALIMENTOS}")

    alimentos, ids, claves = [], set(), set()
    for n, fila in enumerate(filas[1:], start=2):
        fila = (list(fila) + [None] * 13)[:13]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
            continue  # fila totalmente vacía: se ignora
        (vid, nombre, sinonimos, grupo, estado, kcal, prot, carb, gras,
         fibra, fuente, id_fuente, licencia) = fila

        vid = texto(vid, "id", n)
        if not RE_ID.match(vid):
            fallo(f"fila {n}: id '{vid}' no tiene la forma ALI-0000")
        if vid in ids:
            fallo(f"fila {n}: id '{vid}' duplicado")
        ids.add(vid)

        nombre = texto(nombre, "nombre", n)
        grupo = texto(grupo, "grupo", n)
        estado = texto(estado, "estado", n)
        fuente = texto(fuente, "fuente", n)
        licencia = texto(licencia, "licencia", n)
        clave = (normalizar(nombre), estado)
        if clave in claves:
            fallo(f"fila {n}: alimento duplicado (nombre+estado): {nombre} / {estado}")
        claves.add(clave)

        sin_txt = texto(sinonimos, "sinonimos", n, obligatorio=False)
        lista_sin = [s.strip() for s in sin_txt.split(";") if s.strip()] if sin_txt else []

        kcal = numero(kcal, "kcal", n)
        prot = numero(prot, "proteina_g", n)
        carb = numero(carb, "carbohidratos_g", n)
        gras = numero(gras, "grasa_g", n)
        fibra = numero(fibra, "fibra_g", n, obligatorio=False)

        if id_fuente is None or (isinstance(id_fuente, str) and not id_fuente.strip()):
            fallo(f"fila {n}: 'id_fuente' vacío")
        if isinstance(id_fuente, float) and id_fuente.is_integer():
            id_fuente = int(id_fuente)
        if isinstance(id_fuente, str):
            id_fuente = id_fuente.strip()

        atwater = 4 * prot + 4 * carb + 9 * gras
        alimentos.append({
            "id": vid, "nombre": nombre, "sinonimos": lista_sin,
            "grupo": grupo, "estado": estado, "kcal": kcal,
            "proteina_g": prot, "carbohidratos_g": carb, "grasa_g": gras,
            "fibra_g": fibra, "fuente": fuente, "id_fuente": id_fuente,
            "licencia": licencia, "porciones": [],
            "kcalDesvia": abs(kcal - atwater) / max(kcal, 1) > 0.15,
        })

    if not alimentos:
        fallo("la hoja 'Alimentos' no tiene filas")
    por_id = {a["id"]: a for a in alimentos}

    # ---------------- hoja Porciones
    wp = wb["Porciones"]
    pfilas = list(wp.iter_rows(values_only=True))
    if not pfilas or [c if c is not None else "" for c in pfilas[0]][:5] != COLS_PORCIONES:
        fallo(f"encabezados de 'Porciones' distintos de {COLS_PORCIONES}")
    for n, fila in enumerate(pfilas[1:], start=2):
        fila = (list(fila) + [None] * 5)[:5]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
            continue
        vid, _alimento, medida, gramos, fuente = fila
        vid = texto(vid, "id", n)
        if vid not in por_id:
            fallo(f"Porciones fila {n}: id '{vid}' no existe en 'Alimentos'")
        medida = texto(medida, "medida", n)
        fuente = texto(fuente, "fuente", n)
        gramos = numero(gramos, "gramos", n)
        if gramos <= 0:
            fallo(f"Porciones fila {n}: gramos={gramos} debe ser > 0")
        # Contrato del Diario (js/diario.js): la porcion viaja como
        # {nombre, g, fuente}. La hoja conserva sus encabezados medida/gramos;
        # el renombre vive SOLO aqui, en la frontera con la app.
        por_id[vid]["porciones"].append(
            {"nombre": medida, "g": gramos, "fuente": fuente})

    # ---------------- salida
    out = args.out or os.path.join(os.path.dirname(os.path.abspath(args.xlsx)),
                                   "data", "alimentos.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(alimentos, f, ensure_ascii=False, indent=1)
        f.write("\n")
    n_desvia = sum(1 for a in alimentos if a["kcalDesvia"])
    print(f"OK: {len(alimentos)} alimentos validados "
          f"({n_desvia} con kcalDesvia) → {out}")


if __name__ == "__main__":
    main()

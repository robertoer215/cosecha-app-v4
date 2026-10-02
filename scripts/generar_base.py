#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_base.py — Pipeline reproducible de la base de alimentos COSECHA.

PIPELINE COMPLETO (de la descarga al entregable):

  Paso 0 · DESCARGAS (una sola vez; ver datos/build_fdc_sqlite.py)
      Fuente: USDA FoodData Central (https://fdc.nal.usda.gov/, licencia CC0 1.0,
      dominio público). Se descargan los CSV oficiales de los datasets
      Foundation Foods, SR Legacy y FNDDS a datos/descargas/.
      NO se toca ninguna base propietaria (MyFitnessPal, Fitia, FatSecret, WHOOP).

  Paso 1 · SQLITE (datos/build_fdc_sqlite.py → datos/fdc.sqlite)
      Dos tablas:
        candidatos(fdc_id, dataset, descripcion, kcal, prot, carb, gras, fibra)
          — valores por 100 g VERBATIM de los CSV de USDA.
        porciones(fdc_id, descripcion, gramos)
          — porciones caseras (food_portion.csv), gramos VERBATIM.

  Paso 2 · SELECCIONES (datos/curar_*.py, datos/build_*.py, datos/gen_*.py
           → datos/seleccion-<grupo>.json, 9 grupos)
      Cada script de curaduría elige alimentos relevantes para México desde
      candidatos, les pone nombre y sinónimos es-MX y copia los macros VERBATIM
      del sqlite (la traducción solo toca el nombre, nunca una cifra).

  Paso 3 · FUSIÓN (este script)
      a) Carga las 9 seleccion-*.json y re-valida contra fdc.sqlite que cada
         cifra sigue siendo VERBATIM (si una difiere, el script ABORTA).
      b) Deduplica: primero por id_fuente (misma comida USDA elegida por dos
         grupos; gana la primera aparición en orden alfabético de archivo),
         después por clave (nombre normalizado sin acentos + estado) con
         prioridad Foundation > SR Legacy > FNDDS.
      c) Añade la carta COSECHA leyendo (SOLO lectura) js/data.js (ING),
         js/bebidas.js (BEB y BEB_ADDON) y js/postres.js (POS) del repo
         cosecha-app-v4 vía Node. Cada item entra por 100 g derivado de su
         porción Estándar: macros de la etiqueta ÷ gramos × 100, redondeado a
         2 decimales (conversión de unidades, no corrección). Las bebidas solo
         declaran mililitros, así que su fila es por 100 ml (densidad NO
         asumida: la cifra es etiqueta ÷ ml × 100 y así se declara en el Leeme).
      d) Pega hasta 3 porciones caseras por alimento USDA (tabla porciones,
         nombre traducido es-MX, gramos VERBATIM, fuente 'USDA FDC'). Una
         descripción sin traducción segura se EXCLUYE, nunca se estima.
      e) Ordena por grupo y nombre, asigna ids ALI-0001…, marca kcalDesvia
         (|kcal − (4P+4C+9G)| / max(kcal,1) > 0.15) SIN corregir nada.
      f) Escribe data/alimentos.json y COSECHA_Base_Alimentos.xlsx.

  Paso 4 · ROUND-TRIP (scripts/excel_a_json.py)
      Lee el xlsx, valida esquema y tipos y regenera un alimentos.json
      equivalente. Es la vía para editar la base desde Excel.

Uso:
  python3 generar_base.py [--datos DIR] [--repo DIR] [--out DIR]

  --datos  carpeta con seleccion-*.json y fdc.sqlite   (default: ../../datos)
  --repo   repo cosecha-app-v4, SOLO LECTURA           (default: ~/cosecha-app-v4)
  --out    carpeta de salida (staging-datos)           (default: ../)
"""
import argparse
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import unicodedata
from glob import glob

# ---------------------------------------------------------------- utilidades

RANK_DATASET = {"Foundation": 0, "SR Legacy": 1, "FNDDS": 2}

CAMPOS = ["id", "nombre", "sinonimos", "grupo", "estado", "kcal", "proteina_g",
          "carbohidratos_g", "grasa_g", "fibra_g", "fuente", "id_fuente",
          "licencia", "porciones", "kcalDesvia"]


def normalizar(s):
    """minúsculas + sin acentos, para claves de deduplicación y orden."""
    s = unicodedata.normalize("NFD", s.lower().strip())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def kcal_desvia(kcal, p, c, g):
    """|kcal − (4P+4C+9G)| / max(kcal, 1) > 0.15 — solo marca, nunca corrige."""
    atwater = 4 * p + 4 * c + 9 * g
    return abs(kcal - atwater) / max(kcal, 1) > 0.15


# ------------------------------------------------- traducción de porciones

# Descripciones de food_portion que NO son medida casera útil: se excluyen
# (excluir está permitido; estimar no).
PATRONES_EXCLUIR = re.compile(
    r"quantity not specified|guideline amount|nlea serving|cubic inch|"
    r"surface inch|yield|^1 lb$|^1 quart$|100 calorie|\(303 x 406\)",
    re.IGNORECASE)

# Traducciones exactas (descripcion USDA → es-MX). Solo entra lo que tiene
# traducción segura; el resto se excluye.
TRAD_EXACTA = {
    "1 cup": "1 taza", "0.5 cup": "½ taza", "0.25 cup": "¼ taza",
    "0.75 cup": "¾ taza", "2 cup": "2 tazas",
    "1 cup (8 fl oz)": "1 taza (240 ml)",
    "1 cup, sliced": "1 taza en rebanadas", "1 cup slices": "1 taza en rebanadas",
    "0.5 cup slices": "½ taza en rebanadas",
    "1 cup, chopped": "1 taza picada", "1 cup chopped": "1 taza picada",
    "1 cup, chopped or diced": "1 taza picada o en cubos",
    "0.5 cup, chopped or diced": "½ taza picada o en cubos",
    "1 cup, shredded": "1 taza rallada",
    "1 cup, diced": "1 taza en cubos", "1 cup, cubes": "1 taza en cubos",
    "1 cup cubes": "1 taza en cubos", "1 cup pieces": "1 taza en trozos",
    "1 cup, cooked": "1 taza cocida", "1 cup, cooked, diced": "1 taza cocida en cubos",
    "1 cup, halves": "1 taza en mitades", "1 cup halves": "1 taza en mitades",
    "1 cup, whole": "1 taza entera", "1 cup, mashed": "1 taza en puré",
    "1 cup, pulp": "1 taza de pulpa", "1 cup, NFS": "1 taza",
    "1 cup (not packed)": "1 taza (sin compactar)",
    "1 cup, unthawed": "1 taza (congelada)",
    "1 tbsp": "1 cucharada", "1 tablespoon": "1 cucharada",
    "1 tbsp chopped": "1 cucharada picada",
    "1 tsp": "1 cucharadita", "1 teaspoon": "1 cucharadita",
    "1 oz": "1 onza", "2 oz": "2 onzas", "3 oz": "3 onzas", "4 oz": "4 onzas",
    "1 fl oz": "1 onza líquida", "1 fl oz (NFS)": "1 onza líquida",
    "1 fl oz (no ice)": "1 onza líquida (sin hielo)",
    "1 fl oz (with ice)": "1 onza líquida (con hielo)",
    "1 large": "1 pieza grande", "1 medium": "1 pieza mediana",
    "1 small": "1 pieza chica", "1 small/regular": "1 pieza chica/regular",
    "1 regular/large": "1 pieza regular/grande", "1 regular": "1 pieza regular",
    "1 small/individual": "1 pieza chica/individual",
    "1 miniature": "1 miniatura", "1 miniature/bite size": "1 miniatura (bocado)",
    "1 item, any size": "1 pieza (cualquier tamaño)", "1 piece": "1 pieza",
    "1 unit": "1 unidad",
    "1 fruit": "1 fruta", "1 fruit without refuse": "1 fruta (parte comestible)",
    "1 slice": "1 rebanada", "1 slice (1 oz)": "1 rebanada (1 onza)",
    "1 medium or regular slice": "1 rebanada mediana",
    "1 large or thick slice": "1 rebanada gruesa",
    "1 small or thin/very thin slice": "1 rebanada delgada",
    "1 slice, snack-size": "1 rebanada chica",
    "1 slice, crust not eaten": "1 rebanada (sin orilla)",
    "1 slice/chunk": "1 rebanada o trozo",
    "1 fillet": "1 filete", "0.5 fillet": "½ filete",
    "1 large fillet": "1 filete grande", "1 small/regular fillet": "1 filete chico/regular",
    "1 egg": "1 huevo", "1 pepper": "1 chile", "1 leaf": "1 hoja",
    "1 head": "1 cabeza", "1 kernel": "1 grano", "1 chip": "1 totopo/chip",
    "1 can": "1 lata",
    "1 can or bottle (12 fl oz)": "1 lata o botella (355 ml)",
    "1 can or bottle (16 fl oz)": "1 lata o botella (473 ml)",
    "1 steak": "1 bistec", "1 breast": "1 pechuga",
    "1 link": "1 pieza (salchicha)", "1 stick": "1 barrita",
    "1 drumstick with skin": "1 pierna con piel",
    "1 thigh with skin": "1 muslo con piel",
    "1 taco salad, NS as to size": "1 ensalada de taco",
    "1 small taco salad": "1 ensalada de taco chica",
    "1 medium taco salad": "1 ensalada de taco mediana",
    "1 large taco salad": "1 ensalada de taco grande",
    "1 enchilada, any size": "1 enchilada", "1 fajita": "1 fajita",
    "1 taco or tostada": "1 taco o tostada", "1 taco": "1 taco",
    "1 tostada": "1 tostada", "1 tamale": "1 tamal", "1 burrito": "1 burrito",
    "1 quesadilla": "1 quesadilla", "1 tortilla": "1 tortilla",
    "1 order": "1 orden", "1 drink": "1 bebida",
    "1 package": "1 paquete", "1 packet": "1 sobre",
    "1 container (8 oz)": "1 envase (8 onzas)",
    "1 container (6 oz)": "1 envase (6 onzas)",
    "1 individual container": "1 envase individual",
    "1 prepackaged single serving": "1 porción individual empacada",
    "1 serving": "1 porción",
    "1 whole fish, any size": "1 pescado entero",
    "1 half circle": "medio círculo", "1 full circle": "1 círculo completo",
    "1 triangular piece": "1 pieza triangular",
    "1 large single serving bag": "1 bolsa individual grande",
    "1 medium single serving bag": "1 bolsa individual mediana",
    "1 small single serving bag": "1 bolsa individual chica",
    "1 dipping-size container": "1 envase para dip",
    "1 bar": "1 barra", "1 roll": "1 rollo", "1 bun": "1 bollo",
    "1 bagel": "1 bagel", "1 muffin": "1 muffin", "1 biscuit": "1 bisquet",
    "1 pancake": "1 hotcake", "1 waffle": "1 waffle", "1 cookie": "1 galleta",
    "1 donut": "1 dona", "1 shrimp": "1 camarón", "1 wing": "1 ala",
    "1 drumstick": "1 pierna", "1 thigh": "1 muslo", "1 ear": "1 elote (mazorca)",
    "1 clove": "1 diente", "1 stalk": "1 tallo", "1 sprig": "1 ramita",
    "1 wedge": "1 gajo", "1 ring": "1 aro", "1 pod": "1 vaina",
    "1 spear": "1 espárrago (tallo)", "1 floweret": "1 florete",
    "1 floret": "1 florete", "1 tomato": "1 jitomate", "1 potato": "1 papa",
    "1 banana": "1 plátano", "1 apple": "1 manzana", "1 orange": "1 naranja",
    "1 patty": "1 tortita", "1 sausage": "1 salchicha", "1 chop": "1 chuleta",
    "1 cutlet": "1 milanesa", "1 strip": "1 tira", "1 cube": "1 cubo",
    "1 ball": "1 bola", "1 scoop": "1 bola (scoop)",
    "1 bottle": "1 botella", "1 juice box": "1 cajita de jugo",
    "1 pouch": "1 bolsita", "1 glass": "1 vaso",
}

# Orden de preferencia al elegir hasta 3 porciones por alimento.
def _rango_preferencia(medida):
    if "taza" in medida:
        return 0
    if medida.startswith(("1 cucharada", "1 cucharadita")):
        return 2
    if "onza líquida" in medida:
        return 5
    if re.match(r"^\d+ onzas?$", medida) or "onza" in medida:
        return 4
    if medida.startswith(("1 lata", "1 botella", "1 envase", "1 paquete",
                          "1 sobre", "1 porción", "1 bolsa", "1 vaso",
                          "1 cajita", "1 bolsita")):
        return 3
    return 1  # piezas, frutas, rebanadas, filetes, platillos…


def traducir_porciones(filas_porcion):
    """[(descripcion, gramos)] → hasta 3 dicts es-MX; lo no traducible se excluye."""
    out = []
    for desc, gramos in filas_porcion:
        if gramos is None or gramos <= 0:
            continue
        d = desc.strip()
        if PATRONES_EXCLUIR.search(d):
            continue
        medida = TRAD_EXACTA.get(d) or TRAD_EXACTA.get(d.lower())
        if medida is None:
            continue  # sin traducción segura → se excluye, no se estima
        out.append({"medida": medida, "gramos": gramos, "fuente": "USDA FDC"})
    # dedup por nombre de medida (conserva la primera) y elige por preferencia
    vistos, unicos = set(), []
    for p in out:
        if p["medida"] in vistos:
            continue
        vistos.add(p["medida"])
        unicos.append(p)
    unicos.sort(key=lambda p: _rango_preferencia(p["medida"]))
    return unicos[:3]


# ------------------------------------------------------------ carta COSECHA

NODE_SNIPPET = """
const { pathToFileURL } = require('url');
(async () => {
  const base = process.argv[1];
  const d = await import(pathToFileURL(base + '/js/data.js'));
  const b = await import(pathToFileURL(base + '/js/bebidas.js'));
  const p = await import(pathToFileURL(base + '/js/postres.js'));
  process.stdout.write(JSON.stringify({
    ING: d.ING, BEB: b.BEB, ADDON: b.BEB_ADDON, POS: p.POS
  }));
})().catch(e => { console.error(e); process.exit(1); });
"""


def leer_carta(repo):
    """Lee ING/BEB/BEB_ADDON/POS del repo (SOLO lectura) vía Node."""
    r = subprocess.run(["node", "-e", NODE_SNIPPET, repo],
                       capture_output=True, text=True, check=True)
    return json.loads(r.stdout)


def carta_a_filas(carta):
    """Cada item de la carta entra por 100 g (bebidas: por 100 ml) derivado de
    su porción Estándar: etiqueta ÷ gramos × 100, redondeado a 2 decimales.
    Es conversión de unidades, no corrección."""
    filas = []

    def alta(item, base_g, es_ml, porciones):
        por100 = lambda v: round(v / base_g * 100, 2)
        filas.append({
            "nombre": item["nombre"],
            "sinonimos": [],
            "grupo": "Carta COSECHA",
            "estado": "listo para comer",
            "kcal": por100(item["kcal"]),
            "proteina_g": por100(item["prot"]),
            "carbohidratos_g": por100(item["carb"]),
            "grasa_g": por100(item["gras"]),
            "fibra_g": None,  # la etiqueta de la carta no declara fibra
            "fuente": "Carta COSECHA (js/data.js)",
            "id_fuente": item["id"],
            "licencia": "Propia",
            "porciones": porciones,
            "_es_ml": es_ml,
        })

    for it in carta["ING"]:
        alta(it, it["g"], False,
             [{"medida": "1 porción Estándar", "gramos": it["g"],
               "fuente": "Carta COSECHA"}])
    for it in carta["BEB"]:
        # Las bebidas declaran ml: la fila queda por 100 ml (sin asumir densidad);
        # no se les pega porción en gramos porque la carta no declara gramos.
        alta(it, it["ml"], True, [])
    alta(carta["ADDON"], carta["ADDON"]["ml"], True, [])
    for it in carta["POS"]:
        alta(it, it["g"], False,
             [{"medida": "1 porción Estándar", "gramos": it["g"],
               "fuente": "Carta COSECHA"}])
    return filas


# ------------------------------------------------------------------- fusión

def fusionar(dir_datos, repo):
    # --- selecciones USDA
    archivos = sorted(glob(os.path.join(dir_datos, "seleccion-*.json")))
    if len(archivos) != 9:
        sys.exit(f"ERROR: esperaba 9 seleccion-*.json, hay {len(archivos)}")
    filas = []
    for a in archivos:
        with open(a, encoding="utf-8") as f:
            filas.extend(json.load(f))
    n_inicial = len(filas)

    # --- re-validación VERBATIM contra el sqlite (aborta si algo difiere)
    con = sqlite3.connect(os.path.join(dir_datos, "fdc.sqlite"))
    cur = con.cursor()
    for r in filas:
        fila_db = cur.execute(
            "SELECT kcal, prot, carb, gras, fibra, dataset FROM candidatos "
            "WHERE fdc_id=?", (r["id_fuente"],)).fetchone()
        if fila_db is None:
            sys.exit(f"ERROR: fdc_id {r['id_fuente']} no está en fdc.sqlite")
        esperado = (r["kcal"], r["proteina_g"], r["carbohidratos_g"],
                    r["grasa_g"], r["fibra_g"], r["dataset"])
        for a, b in zip(esperado, fila_db):
            if a is None and b is None:
                continue
            if a != b:
                sys.exit(f"ERROR: fdc_id {r['id_fuente']} NO es verbatim: "
                         f"json={esperado} sqlite={fila_db}")

    # --- dedup 1: id_fuente (misma comida USDA en dos grupos; macros idénticos
    #     por construcción — validado arriba). Gana la primera aparición en
    #     orden alfabético de archivo.
    vistos, sin_rep = set(), []
    for r in filas:
        if r["id_fuente"] in vistos:
            continue
        vistos.add(r["id_fuente"])
        sin_rep.append(r)

    # --- dedup 2: clave nombre normalizado + estado; gana
    #     Foundation > SR Legacy > FNDDS (empate: primera aparición).
    mejor, orden = {}, []
    for r in sin_rep:
        k = (normalizar(r["nombre"]), r["estado"])
        if k not in mejor:
            mejor[k] = r
            orden.append(k)
        elif RANK_DATASET[r["dataset"]] < RANK_DATASET[mejor[k]["dataset"]]:
            mejor[k] = r
    usda = [mejor[k] for k in orden]

    # --- porciones caseras USDA (hasta 3, traducidas, gramos verbatim)
    for r in usda:
        pts = cur.execute(
            "SELECT descripcion, gramos FROM porciones WHERE fdc_id=?",
            (r["id_fuente"],)).fetchall()
        r["porciones"] = traducir_porciones(pts)
        r["fuente"] = f"USDA FDC ({r['dataset']})"
        r["licencia"] = "CC0 1.0"
    con.close()

    # --- carta COSECHA
    carta = carta_a_filas(leer_carta(repo))
    for c in carta:
        k = (normalizar(c["nombre"]), c["estado"])
        if k in mejor:
            sys.exit(f"ERROR: la carta duplica un alimento USDA: {c['nombre']}")

    # --- ensamblado final: orden por grupo y nombre, ids ALI-0001…
    todas = usda + carta
    todas.sort(key=lambda r: (normalizar(r["grupo"]), normalizar(r["nombre"])))
    final = []
    for i, r in enumerate(todas, 1):
        final.append({
            "id": f"ALI-{i:04d}",
            "nombre": r["nombre"],
            "sinonimos": r["sinonimos"],
            "grupo": r["grupo"],
            "estado": r["estado"],
            "kcal": r["kcal"],
            "proteina_g": r["proteina_g"],
            "carbohidratos_g": r["carbohidratos_g"],
            "grasa_g": r["grasa_g"],
            "fibra_g": r["fibra_g"],
            "fuente": r["fuente"],
            "id_fuente": r["id_fuente"],
            "licencia": r["licencia"],
            "porciones": r["porciones"],
            "kcalDesvia": kcal_desvia(r["kcal"], r["proteina_g"],
                                      r["carbohidratos_g"], r["grasa_g"]),
        })

    # --- criterios de aceptación
    assert len(final) >= 1500, f"solo {len(final)} filas"
    claves = {(normalizar(r["nombre"]), r["estado"]) for r in final}
    assert len(claves) == len(final), "quedan duplicados"
    for r in final:
        assert r["fuente"] and r["id_fuente"] is not None and r["licencia"]
        for m in ("kcal", "proteina_g", "carbohidratos_g", "grasa_g"):
            assert isinstance(r[m], (int, float)), f"{r['id']} macro vacío: {m}"
    print(f"filas iniciales USDA: {n_inicial} → sin repetir id_fuente: "
          f"{len(sin_rep)} → sin repetir clave: {len(usda)} | carta: "
          f"{len(carta)} | TOTAL: {len(final)}")
    return final


# -------------------------------------------------------------------- salida

def escribir_json(final, ruta):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8") as f:
        json.dump(final, f, ensure_ascii=False, indent=1)
        f.write("\n")


LEEME = [
    "COSECHA — Base de alimentos (cómo editar esta hoja)",
    "",
    "1. Edita SOLO las hojas 'Alimentos' y 'Porciones'. No cambies los",
    "   encabezados ni el orden de las columnas.",
    "2. 'Alimentos': una fila por alimento. kcal y macros son POR 100 g",
    "   (en los items de bebida de la Carta COSECHA, por 100 ml: la carta",
    "   declara mililitros y aquí no se asume densidad). 'sinonimos' es una",
    "   lista separada por '; '. 'fibra_g' puede quedar vacía (dato no",
    "   declarado por la fuente); kcal, proteina_g, carbohidratos_g y",
    "   grasa_g NUNCA pueden quedar vacíos.",
    "3. Todo alimento lleva fuente, id_fuente y licencia. Un valor",
    "   nutricional debe ser VERBATIM de su fuente: si no tienes el dato,",
    "   borra la fila o deja fibra_g vacía; jamás lo estimes.",
    "4. 'Porciones': medidas caseras por alimento (máx. 3 recomendado).",
    "   'id' debe existir en 'Alimentos'; 'gramos' es número > 0.",
    "5. Para regenerar el JSON de la app ejecuta:",
    "       python3 scripts/excel_a_json.py COSECHA_Base_Alimentos.xlsx",
    "   Valida esquema y tipos y escribe data/alimentos.json (misma",
    "   estructura que el original; el campo kcalDesvia se recalcula solo:",
    "   marca |kcal − (4P+4C+9G)| / max(kcal,1) > 0.15 y NO corrige nada).",
    "6. Fuentes y licencias: hoja 'Fuentes'. USDA FDC es CC0 1.0 (dominio",
    "   público); la Carta COSECHA es propia. No se usó ninguna base",
    "   propietaria (MyFitnessPal, Fitia, FatSecret, WHOOP).",
]

FUENTES = [
    ("USDA FDC (Foundation)", "CC0 1.0 (dominio público)",
     "U.S. Department of Agriculture, Agricultural Research Service. "
     "FoodData Central: Foundation Foods.", "https://fdc.nal.usda.gov/"),
    ("USDA FDC (SR Legacy)", "CC0 1.0 (dominio público)",
     "U.S. Department of Agriculture, Agricultural Research Service. "
     "FoodData Central: SR Legacy.", "https://fdc.nal.usda.gov/"),
    ("USDA FDC (FNDDS)", "CC0 1.0 (dominio público)",
     "U.S. Department of Agriculture, Agricultural Research Service. "
     "FoodData Central: FNDDS 2021-2023.", "https://fdc.nal.usda.gov/"),
    ("USDA FDC", "CC0 1.0 (dominio público)",
     "Porciones caseras: USDA FoodData Central, food_portion (gramos "
     "verbatim; nombre de la medida traducido al es-MX).",
     "https://fdc.nal.usda.gov/"),
    ("Carta COSECHA (js/data.js)", "Propia",
     "Carta propia de COSECHA. Items y etiquetas de los archivos js/data.js "
     "(ING), js/bebidas.js (BEB, BEB_ADDON) y js/postres.js (POS) del repo "
     "cosecha-app-v4; cada item convertido a por-100 g (bebidas: por-100 ml) "
     "desde su porción Estándar.", ""),
    ("Carta COSECHA", "Propia",
     "Porción 'Estándar' de la carta (gramos verbatim de js/data.js y "
     "js/postres.js).", ""),
]

COLS_ALIMENTOS = ["id", "nombre", "sinonimos", "grupo", "estado", "kcal",
                  "proteina_g", "carbohidratos_g", "grasa_g", "fibra_g",
                  "fuente", "id_fuente", "licencia"]
COLS_PORCIONES = ["id", "alimento", "medida", "gramos", "fuente"]
COLS_FUENTES = ["fuente", "licencia", "atribución", "url"]


def escribir_xlsx(final, ruta):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    negrita = Font(bold=True)

    ws = wb.active
    ws.title = "Alimentos"
    ws.append(COLS_ALIMENTOS)
    for c in ws[1]:
        c.font = negrita
    for r in final:
        ws.append([r["id"], r["nombre"], "; ".join(r["sinonimos"]),
                   r["grupo"], r["estado"], r["kcal"], r["proteina_g"],
                   r["carbohidratos_g"], r["grasa_g"], r["fibra_g"],
                   r["fuente"], r["id_fuente"], r["licencia"]])

    wp = wb.create_sheet("Porciones")
    wp.append(COLS_PORCIONES)
    for c in wp[1]:
        c.font = negrita
    for r in final:
        for p in r["porciones"]:
            wp.append([r["id"], r["nombre"], p["medida"], p["gramos"],
                       p["fuente"]])

    wf = wb.create_sheet("Fuentes")
    wf.append(COLS_FUENTES)
    for c in wf[1]:
        c.font = negrita
    for fila in FUENTES:
        wf.append(list(fila))

    wl = wb.create_sheet("Leeme")
    for i, linea in enumerate(LEEME, 1):
        wl.cell(row=i, column=1, value=linea)
    wl.cell(row=1, column=1).font = negrita

    wb.save(ruta)


def main():
    aqui = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--datos", default=os.path.normpath(
        os.path.join(aqui, "..", "..", "datos")))
    ap.add_argument("--repo", default=os.path.expanduser("~/cosecha-app-v4"))
    ap.add_argument("--out", default=os.path.normpath(os.path.join(aqui, "..")))
    args = ap.parse_args()

    final = fusionar(args.datos, args.repo)
    ruta_json = os.path.join(args.out, "data", "alimentos.json")
    escribir_json(final, ruta_json)
    ruta_xlsx = os.path.join(args.out, "COSECHA_Base_Alimentos.xlsx")
    escribir_xlsx(final, ruta_xlsx)

    n_desvia = sum(1 for r in final if r["kcalDesvia"])
    por_grupo = {}
    for r in final:
        por_grupo[r["grupo"]] = por_grupo.get(r["grupo"], 0) + 1
    con_porcion = sum(1 for r in final if r["porciones"])
    print(f"kcalDesvia=true: {n_desvia} | con porciones caseras: {con_porcion}")
    for g in sorted(por_grupo, key=normalizar):
        print(f"  {por_grupo[g]:5d}  {g}")
    print(f"OK → {ruta_json}\nOK → {ruta_xlsx}")


if __name__ == "__main__":
    main()
